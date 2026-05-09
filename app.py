import os
import threading
import uuid
import fitz  # للـ PDF
from docx import Document  # للـ Word
from pptx import Presentation  # للـ PowerPoint
from pdf2docx import Converter
from gtts import gTTS
from flask import Flask
from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import Application, CommandHandler, MessageHandler, filters, ContextTypes, CallbackQueryHandler

# --- إعداد السيرفر ---
app = Flask(__name__)
@app.route('/')
def home(): return "Tools Bot is Live & Monitored!"

# --- المتغيرات الآمنة ---
BOT_TOKEN = os.environ.get("BOT_TOKEN")
ADMIN_ID = os.environ.get("ADMIN_ID") 
DOWNLOAD_DIR = "downloads"
os.makedirs(DOWNLOAD_DIR, exist_ok=True)

# --- نظام المراقبة التفصيلي (زي الكود القديم) ---
async def notify_admin(context, message):
    if ADMIN_ID:
        try:
            # استخدام parse_mode='Markdown' عشان الرسائل تظهر منسقة
            await context.bot.send_message(chat_id=ADMIN_ID, text=f"📢 *تقرير المراقبة:*\n{message}", parse_mode='Markdown')
        except: pass

# --- وظيفة استخراج النص ---
def get_text_from_any(file_path):
    ext = file_path.lower()
    text = ""
    try:
        if ext.endswith('.pdf'):
            doc = fitz.open(file_path)
            text = "".join([page.get_text() for page in doc])
            doc.close()
        elif ext.endswith('.docx'):
            doc = Document(file_path)
            text = "\n".join([para.text for para in doc.paragraphs])
        elif ext.endswith(('.pptx', '.ppt')):
            prs = Presentation(file_path)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"): text += shape.text + " "
    except: pass
    return text

# --- أوامر البوت ---
async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    user_name = user.first_name
    
    # تقرير تفصيلي للأدمن
    admin_msg = (f"👤 *مستخدم جديد دخل البوت:*\n"
                 f"🔹 الاسم: {user_name}\n"
                 f"🔹 اليوزر: @{user.username if user.username else 'لا يوجد'}\n"
                 f"🔹 الـ ID: `{user.id}`")
    await notify_admin(context, admin_msg)
    
    await update.message.reply_text(f"مرحباً يا {user_name}! 👋\nأرسل ملف (PDF, Word, PPT) وسأعالجه لك فوراً.")

async def handle_document(update: Update, context: ContextTypes.DEFAULT_TYPE):
    doc = update.message.document
    user = update.effective_user
    file_ext = doc.file_name.lower().split('.')[-1]
    
    if file_ext not in ['pdf', 'docx', 'pptx', 'ppt']:
        await update.message.reply_text("عذراً، أقبل ملفات PDF, Word, PowerPoint فقط.")
        return

    # إخطار الإدارة بالملف المرفوع
    await notify_admin(context, f"📥 *{user.first_name}* أرسل ملفاً:\n📄 `{doc.file_name}`")

    file_path = os.path.join(DOWNLOAD_DIR, f"{uuid.uuid4()}_{doc.file_name}")
    file = await context.bot.get_file(doc.file_id)
    await file.download_to_drive(file_path)
    context.user_data["current_file"] = file_path

    keyboard = [
        [InlineKeyboardButton("تحويل لصوت 🎙️", callback_data="audio")],
        [InlineKeyboardButton("تحويل لـ Word 📝", callback_data="word")]
    ]
    await update.message.reply_text(f"📄 الملف جاهز! ماذا تريد أن أفعل؟", reply_markup=InlineKeyboardMarkup(keyboard))

async def button_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    user = query.from_user
    file_path = context.user_data.get("current_file")
    
    if not file_path or not os.path.exists(file_path): return

    # مراقبة ضغطات الأزرار
    await notify_admin(context, f"⚙️ *{user.first_name}* اختار عملية: `{query.data}`")

    if query.data == "audio":
        await query.edit_message_text("🎧 جاري تحويل النص لصوت...")
        text = get_text_from_any(file_path)
        if text.strip():
            audio_path = f"{uuid.uuid4()}.mp3"
            # ضبط الصوت على الإنجليزية en كما طلبت
            tts = gTTS(text=text[:3000], lang='en', slow=False)
            tts.save(audio_path)
            with open(audio_path, "rb") as audio:
                await query.message.reply_audio(audio)
            os.remove(audio_path)
            await notify_admin(context, f"✅ تم تحويل الصوت لـ *{user.first_name}* بنجاح.")
        else:
            await query.message.reply_text("لم أجد نصاً في الملف.")

    elif query.data == "word":
        if not file_path.lower().endswith('.pdf'):
            await query.message.reply_text("هذه الميزة لملفات PDF فقط.")
            return
        await query.edit_message_text("🔄 جاري التحويل لوورد...")
        try:
            docx_path = file_path.replace('.pdf', '.docx')
            cv = Converter(file_path)
            cv.convert(docx_path); cv.close()
            with open(docx_path, "rb") as f:
                await query.message.reply_document(f)
            os.remove(docx_path)
            await notify_admin(context, f"✅ تم تحويل PDF لـ Word لـ *{user.first_name}*.")
        except Exception as e:
            await query.message.reply_text(f"خطأ في التحويل: {e}")

def main():
    if not BOT_TOKEN: return
    threading.Thread(target=lambda: app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 10000))), daemon=True).start()
    application = Application.builder().token(BOT_TOKEN).build()
    application.add_handler(CommandHandler("start", start))
    application.add_handler(MessageHandler(filters.Document.ALL, handle_document))
    application.add_handler(CallbackQueryHandler(button_callback))
    application.run_polling(drop_pending_updates=True)

if __name__ == "__main__":
    main()

